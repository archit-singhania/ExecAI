"""Product behavior acceptance: isolation, persistence, roles and durable board runs."""
import io
import time
import uuid
from fastapi.testclient import TestClient
from app.main import app
from app.database import SessionLocal
from app.models import User, Job
from app.studio_models import StudioRecord, RunEvent
from app.llm import _client as REAL_LLM_CLIENT

client = TestClient(app)

def test_ai_proposals_require_review_and_validate_without_changing_saved_work(monkeypatch):
    import json
    from app import llm_router
    h, _, _ = account(); other, _, _ = account(); sid = workspace(h)
    path = f"/api/studio/{sid}/premortems/proposals"
    monkeypatch.setattr(llm_router, "complete", lambda *a, **k: (json.dumps({"risks": [{"failure": "Supplier delay", "warning": "Missed milestones", "mitigation": "Validate backup"}]}), "fixture-model"))
    assert client.post(path, headers=other, json={"strategy": "Launch"}).status_code == 404
    reply = client.post(path, headers=h, json={"strategy": "Launch"})
    assert reply.status_code == 200 and reply.json()["risks"][0]["origin"] == "ai_proposal"
    assert client.get(f"/api/studio/{sid}/premortems", headers=h).json() == []
    monkeypatch.setattr(llm_router, "complete", lambda *a, **k: ('{"risks":[{"failure":"x","likelihood":"certain"}]}', "fixture-model"))
    assert client.post(path, headers=h, json={"strategy": "Launch"}).status_code == 502
    monkeypatch.setattr(llm_router, "complete", lambda *a, **k: None)
    assert client.post(path, headers=h, json={"strategy": "Launch"}).status_code == 503

def test_practice_transcripts_are_owned_versioned_and_fail_without_saving(monkeypatch):
    from app import llm_router
    h, _, _ = account(); other, _, _ = account(); sid = workspace(h)
    root = f"/api/studio/{sid}/negotiations"
    row = client.post(root, headers=h, json={"title": "Terms", "objectives": "Stay under budget", "constraints": [{"name": "Cost", "maximum": 100}], "packages": [{"name": "A", "values": {"Cost": 90}}]}).json()
    path = root + "/" + row["id"] + "/practice"
    monkeypatch.setattr(llm_router, "complete", lambda *a, **k: ('{"reply":"Which concessions can you trade?"}', "fixture-model"))
    payload = {"message": "Can we discuss price?", "version": row["version"]}
    assert client.post(path, headers=other, json=payload).status_code == 404
    assert client.post(path, headers=h, json={**payload, "version": 999}).status_code == 409
    result = client.post(path, headers=h, json=payload)
    assert result.status_code == 201 and result.json()["data"]["negotiation_version"] == row["version"]
    assert len(client.get(path, headers=h).json()) == 1
    monkeypatch.setattr(llm_router, "complete", lambda *a, **k: ('{"reply":false}', "fixture-model"))
    assert client.post(path, headers=h, json=payload).status_code == 502
    monkeypatch.setattr(llm_router, "complete", lambda *a, **k: None)
    assert client.post(path, headers=h, json=payload).status_code == 503
    assert len(client.get(path, headers=h).json()) == 1
    assert client.get(root, headers=h).json()[0]["data"]["comparison"][0]["within_limits"]

def test_periodic_freshness_catches_due_reviews_and_survives_repeated_ticks():
    from app.freshness_worker import tick
    h, _, _ = account(); sid = workspace(h)
    root = f"/api/studio/{sid}"
    decision = client.post(root + "/records", headers=h, json={"kind": "decision", "title": "Review"}).json()
    research = client.post(root + "/records", headers=h, json={"kind": "research", "title": "Source", "body": "Entered research"}).json()
    assert client.post(root + "/freshness/baselines", headers=h, json={"confirmed": True, "decision_id": decision["id"], "evidence_ids": [research["id"]], "review_dates": {research["id"]: "2000-01-01"}}).status_code == 201
    tick()
    alerts = client.get(root + "/freshness", headers=h).json()["alerts"]
    assert len(alerts) == 1 and alerts[0]["data"]["reason"] == "review_due"
    tick()
    assert len(client.get(root + "/freshness", headers=h).json()["alerts"]) == 1

def test_board_deck_exports_real_files_and_revokes_public_notes_free_snapshots():
    import zipfile
    from pypdf import PdfReader
    from pptx import Presentation
    owner, _, _ = account(); outsider, _, _ = account(); sid = workspace(owner)
    records = f"/api/studio/{sid}/records"; path = f"/api/studio/{sid}/presentations"
    decision = client.post(records, headers=owner, json={"kind": "decision", "title": "Buyer validation", "body": "Interview before expansion", "data": {"alternatives": "Wait", "assumptions": "Demand unknown"}}).json()
    scenario = client.post(records, headers=owner, json={"kind": "scenario", "title": "Conservative", "data": {"cash": 1000, "monthly_cost": 100, "monthly_revenue": 50}}).json()
    result = client.post(path, headers=owner, json={"decision_id": decision["id"], "scenario_ids": [scenario["id"]]})
    assert result.status_code == 201, result.text
    deck = result.json()
    assert client.get(path, headers=outsider).status_code == 404
    slides = deck["data"]["slides"]
    slides[0]["notes"] = "PRIVATE_SPEAKER_SECRET"
    payload = {"title": deck["title"], "slides": slides, "version": deck["version"]}
    saved = client.put(path + "/" + deck["id"], headers=owner, json=payload)
    assert saved.status_code == 200
    assert client.put(path + "/" + deck["id"], headers=owner, json=payload).status_code == 409
    pdf = client.get(f"{path}/{deck['id']}/export/pdf", headers=owner)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert "Interview before expansion" in "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    pptx = client.get(f"{path}/{deck['id']}/export/pptx", headers=owner)
    assert pptx.status_code == 200
    presentation = Presentation(io.BytesIO(pptx.content))
    assert len(presentation.slides) == len(slides) + 1
    assert any(shape.has_chart for slide in presentation.slides for shape in slide.shapes)
    with zipfile.ZipFile(io.BytesIO(pptx.content)) as archive:
        assert all(b"PRIVATE_SPEAKER_SECRET" not in archive.read(name) for name in archive.namelist())
    shared = client.post(f"{path}/{deck['id']}/share", headers=owner)
    assert shared.status_code == 200, shared.text
    slug = shared.json()["slug"]
    public = client.get("/api/share/" + slug)
    assert public.status_code == 200 and public.json()["presentation"]["slides"][0]["body"] == "Interview before expansion"
    assert "PRIVATE_SPEAKER_SECRET" not in public.text
    assert client.delete(f"{path}/{deck['id']}/share", headers=owner).status_code == 200
    assert client.get("/api/share/" + slug).status_code == 404

def test_deck_long_text_exports_paginate_without_losing_content():
    import os
    from pathlib import Path
    from pypdf import PdfReader
    from pptx import Presentation
    from app.presentations import export_slides
    h, _, _ = account(); sid = workspace(h)
    decision = client.post(f"/api/studio/{sid}/records", headers=h, json={"kind": "decision", "title": "Long-form board review", "body": "Initial recommendation"}).json()
    root = f"/api/studio/{sid}/presentations"
    deck = client.post(root, headers=h, json={"decision_id": decision["id"]}).json()
    body = "\n".join(f"Section {i}: validate adoption, costs and delivery before committing funds." for i in range(25))
    slides = [{"title": "A long recommendation " * 6, "body": body, "notes": "PRIVATE_NOTES " * 100}, {"title": "Empty content is permitted", "body": "", "notes": ""}]
    result = client.put(root + "/" + deck["id"], headers=h, json={"title": "Export limits", "slides": slides, "version": deck["version"]})
    assert result.status_code == 200, result.text
    pages = list(export_slides(slides))
    assert len(pages) == 3 and all(len(p["body"].split("\n")) <= 18 for p in pages)
    for format in ("pdf", "pptx"):
        response = client.get(f"{root}/{deck['id']}/export/{format}", headers=h)
        assert response.status_code == 200
        if format == "pdf":
            parsed = PdfReader(io.BytesIO(response.content)); assert len(parsed.pages) == 3
            text = "\n".join(p.extract_text() for p in parsed.pages)
        else:
            parsed = Presentation(io.BytesIO(response.content)); assert len(parsed.slides) == 3
            text = "\n".join(s.text for slide in parsed.slides for s in slide.shapes if s.has_text_frame)
        assert "Section 0:" in text and "Section 24:" in text and "PRIVATE_NOTES" not in text
        if os.getenv("EXEC_AI_QA_OUTPUT"):
            directory = Path(os.environ["EXEC_AI_QA_OUTPUT"]); directory.mkdir(parents=True, exist_ok=True)
            (directory / ("long-board." + format)).write_bytes(response.content)

def test_negotiation_limits_missing_values_and_versioned_readback():
    owner, _, _ = account(); other, _, _ = account(); sid = workspace(owner)
    path = f"/api/studio/{sid}/negotiations"
    payload = {"title": "Supplier terms", "objectives": "Control price and delivery", "alternative": "Use existing supplier", "constraints": [{"name": "Price", "unit": "USD", "maximum": 100}, {"name": "Delivery", "unit": "days", "maximum": 10}], "packages": [{"name": "A", "values": {"Price": 90, "Delivery": 8}}, {"name": "B", "values": {"Price": 105, "Delivery": 7}}]}
    assert client.post(path, headers=other, json=payload).status_code == 404
    result = client.post(path, headers=owner, json=payload)
    assert result.status_code == 201, result.text
    saved = result.json()
    assert [p["within_limits"] for p in saved["data"]["comparison"]] == [True, False]
    assert saved["data"]["comparison"][1]["checks"][0]["violated"]
    payload["version"] = saved["version"]
    payload["packages"][1]["values"]["Price"] = 100
    updated = client.put(path + "/" + saved["id"], headers=owner, json=payload)
    assert updated.status_code == 200 and updated.json()["data"]["comparison"][1]["within_limits"]
    assert client.put(path + "/" + saved["id"], headers=owner, json=payload).status_code == 409
    payload["packages"][0]["values"].pop("Delivery")
    assert client.post(path, headers=owner, json=payload).status_code == 422
    assert len(client.get(path, headers=owner).json()) == 1

def test_freshness_baselines_compare_real_metrics_and_deduplicate_alerts():
    owner, _, _ = account(); stranger, _, _ = account(); sid = workspace(owner)
    records = f"/api/studio/{sid}/records"; path = f"/api/studio/{sid}/freshness"
    decision = client.post(records, headers=owner, json={"kind": "decision", "title": "Scale carefully"}).json()
    metric_payload = {"kind": "metric", "title": "Demand", "data": {"value": 2, "target": 10}}
    metric = client.post(records, headers=owner, json=metric_payload).json()
    baseline = {"confirmed": True, "decision_id": decision["id"], "evidence_ids": [metric["id"]]}
    assert client.post(path + "/baselines", headers=stranger, json=baseline).status_code == 404
    assert client.post(path + "/baselines", headers=owner, json={**baseline, "confirmed": False}).status_code == 422
    assert client.post(path + "/baselines", headers=owner, json=baseline).status_code == 201
    assert client.post(path + "/scan", headers=owner).json()["created"] == 0
    metric_payload.update(version=metric["version"], data={"value": 3, "target": 10})
    assert client.put(records + "/" + metric["id"], headers=owner, json=metric_payload).status_code == 200
    state = client.get(path, headers=owner).json()
    assert len(state["alerts"]) == 1
    alert = state["alerts"][0]
    assert alert["data"]["before"]["data"]["value"] == 2
    assert alert["data"]["after"]["data"]["value"] == 3
    assert client.post(path + "/scan", headers=owner).json()["created"] == 0
    assert client.patch(path + "/alerts/" + alert["id"], headers=owner, json={"state": "acknowledged"}).status_code == 200
    assert client.post(path + "/scan", headers=owner).json()["created"] == 0
    client.delete(records + "/" + metric["id"], headers=owner)
    assert client.post(path + "/scan", headers=owner).json()["created"] == 1
    assert any(a["data"]["reason"] == "removed" for a in client.get(path, headers=owner).json()["alerts"])

def test_research_review_dates_do_not_claim_external_changes():
    owner, _, _ = account(); viewer, _, email = account(); sid = workspace(owner)
    records = f"/api/studio/{sid}/records"; path = f"/api/studio/{sid}/freshness"
    decision = client.post(records, headers=owner, json={"kind": "decision", "title": "Review research"}).json()
    source = client.post(records, headers=owner, json={"kind": "research", "title": "Manual source", "body": "Saved excerpt"}).json()
    baseline = {"confirmed": True, "decision_id": decision["id"], "evidence_ids": [source["id"]], "review_dates": {source["id"]: "2020-01-01"}}
    client.post(path + "/baselines", headers=owner, json=baseline)
    assert client.post(path + "/scan", headers=owner).json()["created"] == 1
    alert = client.get(path, headers=owner).json()["alerts"][0]
    assert alert["data"]["reason"] == "review_due"
    assert "no claim" in alert["data"]["scope"]
    client.post(f"/api/studio/{sid}/members", headers=owner, json={"email": email, "role": "viewer"})
    assert client.get(path, headers=viewer).status_code == 200
    assert client.post(path + "/scan", headers=viewer).status_code == 403

def test_premortem_versioning_and_explicit_idempotent_task_conversion():
    owner, _, _ = account(); other, _, _ = account(); sid = workspace(owner)
    path = f"/api/studio/{sid}/premortems"
    payload = {"title": "Launch risk", "strategy": "Validate before scaling", "risks": [{"failure": "No demand", "impact": "high", "mitigation": "Interview ten buyers", "warning": "No trial conversion"}]}
    created = client.post(path, headers=owner, json=payload)
    assert created.status_code == 201, created.text
    row = created.json(); risk = row["data"]["risks"][0]
    assert client.get(path, headers=other).status_code == 404
    task_path = f"{path}/{row['id']}/risks/{risk['id']}/task"
    assert client.post(task_path, headers=owner, json={"confirmed": False}).status_code == 422
    task = client.post(task_path, headers=owner, json={"confirmed": True}).json()
    assert task["created"] is True
    again = client.post(task_path, headers=owner, json={"confirmed": True}).json()
    assert again == {"task_id": task["task_id"], "created": False}
    payload["version"] = row["version"]
    assert client.put(path + "/" + row["id"], headers=owner, json=payload).status_code == 409
    updated = client.get(path, headers=owner).json()[0]
    payload.update(version=updated["version"], risks=updated["data"]["risks"])
    payload["risks"][0]["warning"] = "Updated warning"
    saved = client.put(path + "/" + row["id"], headers=owner, json=payload)
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["risks"][0]["task_id"] == task["task_id"]
    tasks = client.get(f"/api/studio/{sid}/tasks", headers=owner).json()
    assert len(tasks) == 1 and tasks[0]["priority"] == "High"

def test_decision_graph_rejects_cycles_preserves_records_and_reloads():
    owner, _, _ = account(); sid = workspace(owner)
    path = f"/api/studio/{sid}/decision-graph"
    nodes = [client.post(path + "/assumptions", headers=owner, json={"title": title}).json() for title in ["Demand", "Retention", "Expansion"]]
    edges = []
    for a, b in [(0, 1), (1, 2)]:
        result = client.post(path + "/edges", headers=owner, json={"source": nodes[a]["id"], "target": nodes[b]["id"]})
        assert result.status_code == 201, result.text
        edges.append(result.json())
    closing = {"source": nodes[2]["id"], "target": nodes[0]["id"]}
    assert client.post(path + "/edges", headers=owner, json=closing).status_code == 409
    closing["relationship"] = "evidence"
    assert client.post(path + "/edges", headers=owner, json=closing).status_code == 201
    duplicate = client.post(path + "/edges", headers=owner, json={"source": nodes[0]["id"], "target": nodes[1]["id"]}).json()
    assert duplicate["id"] == edges[0]["id"]
    value = client.get(path, headers=owner, params={"changed": nodes[0]["id"]}).json()
    assert set(value["affected"]) == {nodes[1]["id"], nodes[2]["id"]}
    assert all(n["version"] == 1 for n in value["nodes"])
    assert client.delete(path + "/edges/" + edges[0]["id"], headers=owner).status_code == 200
    assert not client.get(path, headers=owner, params={"changed": nodes[0]["id"]}).json()["affected"]

def test_decision_graph_enforces_workspace_and_viewer_boundaries():
    owner, _, _ = account(); stranger, _, email = account()
    sid = workspace(owner); foreign = workspace(stranger)
    path = f"/api/studio/{sid}/decision-graph"
    a = client.post(path + "/assumptions", headers=owner, json={"title": "Owned"}).json()
    b = client.post(f"/api/studio/{foreign}/decision-graph/assumptions", headers=stranger, json={"title": "Private"}).json()
    assert client.get(path, headers=stranger).status_code == 404
    assert client.post(path + "/edges", headers=owner, json={"source": a["id"], "target": b["id"]}).status_code == 404
    client.post(f"/api/studio/{sid}/members", headers=owner, json={"email": email, "role": "viewer"})
    assert client.get(path, headers=stranger).status_code == 200
    assert client.post(path + "/assumptions", headers=stranger, json={"title": "Denied"}).status_code == 403
    assert client.post(path + "/edges", headers=stranger, json={"source": a["id"], "target": a["id"]}).status_code == 403
    assert client.post(path + "/assumptions", headers=owner, json={"title": "  "}).status_code == 422

def test_reading_review_settings_does_not_opt_in_to_scheduled_email():
    owner, _, _ = account()
    value = client.get("/api/review-schedule", headers=owner).json()
    assert value["cadence"] == "off" and not value["email_enabled"]

def test_legacy_client_obeys_local_only_before_constructing_hosted_client(monkeypatch):
    from app import llm, llm_router
    from types import SimpleNamespace
    settings = SimpleNamespace(llm_provider="groq", llm_local_only=False, groq_api_key="configured", groq_model="test")
    monkeypatch.setattr(llm, "get_settings", lambda: settings)
    # Capture the real function at collection; all normal tests keep providers isolated.
    real_client = REAL_LLM_CLIENT
    token = llm_router.RUN_OPTIONS.set({"local_only": True})
    try:
        assert real_client() == (None, None)
    finally:
        llm_router.RUN_OPTIONS.reset(token)

def test_shared_viewer_cannot_change_tasks_through_legacy_route():
    owner, _, _ = account(); viewer, _, email = account(); sid = workspace(owner)
    task = client.post(f"/api/studio/{sid}/tasks", headers=owner, json={"title": "Owned task"}).json()
    client.post(f"/api/studio/{sid}/members", headers=owner, json={"email": email, "role": "viewer"})
    assert client.patch(f"/api/tasks/{task['id']}", headers=viewer, json={"status": "Done"}).status_code == 403

def test_specialist_debate_is_scoped_and_preserves_two_sources():
    owner, _, _ = account(); other, _, _ = account(); sid = workspace(owner)
    path = f"/api/studio/{sid}/debate"
    payload = {"question": "Should we launch now or validate demand?", "specialist": "cfo", "challenger": "product"}
    assert client.post(path, headers=other, json=payload).status_code == 404
    value = client.post(path, headers=owner, json=payload)
    assert value.status_code == 200, value.text
    reports = value.json()["data"]["reports"]
    assert len(reports) == 2 and reports[0]["agent"] != reports[1]["agent"]
    assert all(r["source"] == "local-template" for r in reports)
    payload["challenger"] = "cfo"
    assert client.post(path, headers=owner, json=payload).status_code == 422

def test_metric_history_and_recipient_only_mentions():
    owner, _, _ = account(); colleague, user, email = account(); sid = workspace(owner)
    client.post(f"/api/studio/{sid}/members", headers=owner, json={"email": email, "role": "editor"})
    path = f"/api/studio/{sid}/records"
    payload = {"kind": "metric", "title": "Buyer interviews", "data": {"value": 2, "target": 5, "objective": "Validate demand"}}
    metric = client.post(path, headers=owner, json=payload).json()
    payload.update(version=metric["version"], data={"value": 5, "target": 5, "objective": "Validate demand"})
    assert client.put(path + "/" + metric["id"], headers=owner, json=payload).status_code == 200
    history = client.get(path, headers=owner).json()
    assert any(r["kind"] == "metric_observation" and r["data"]["value"] == 2 for r in history)
    client.post(path, headers=owner, json={"kind": "comment", "title": "Review the evidence", "body": "@" + email + " Please inspect these interviews."})
    notices = client.get(f"/api/studio/{sid}/notifications", headers=colleague).json()
    assert notices and notices[0]["data"]["recipient_id"] == user["id"]
    assert not any(r["kind"] == "notification" for r in client.get(path, headers=owner).json())
    assert client.post(f"/api/studio/{sid}/notifications/{notices[0]['id']}/read", headers=owner).status_code == 404
    assert client.post(f"/api/studio/{sid}/notifications/{notices[0]['id']}/read", headers=colleague).status_code == 200

def test_reusable_briefs_and_archive_preserve_records():
    owner, _, _ = account(); sid = workspace(owner); path = f"/api/studio/{sid}/records"
    brief = client.post(path, headers=owner, json={"kind": "template", "title": "Launch gate", "body": "Assess launch evidence"}).json()
    archive = client.post(path, headers=owner, json={"kind": "archive", "title": "Archive"}).json()
    assert next(w for w in client.get("/api/studio/workspaces", headers=owner).json() if w["id"] == sid)["archived"]
    client.delete(path + "/" + archive["id"], headers=owner)
    assert any(r["id"] == brief["id"] for r in client.get(path, headers=owner).json())

def test_timezone_schedule_handles_daylight_saving_transition():
    from datetime import datetime
    from app.scheduling import compute_next_run
    before = compute_next_run("weekly", 0, 9, 0, now=datetime(2026, 3, 2, 15), timezone="America/New_York")
    assert before == datetime(2026, 3, 9, 13)  # 09:00 EDT after the spring transition.
    owner, _, _ = account()
    result = client.put("/api/review-schedule", headers=owner, json={"cadence": "weekly", "weekday": 0, "hour": 9, "timezone": "America/New_York"})
    assert result.status_code == 200 and result.json()["timezone"] == "America/New_York"
    assert client.put("/api/review-schedule", headers=owner, json={"timezone": "Mars/Nope"}).status_code == 422

def test_cookie_mutation_rejects_untrusted_origin():
    owner, _, _ = account()
    result = client.post("/api/sessions", headers={**owner, "Origin": "https://untrusted.example"}, json={"business_goal": "Inspect a company"})
    assert result.status_code == 403

def account():
    email = f"{uuid.uuid4().hex}@example.com"
    result = client.post("/api/auth/signup", json={"name": "Founder", "email": email, "password": "correct-horse-battery"})
    assert result.status_code == 200, result.text
    body = result.json()
    with SessionLocal() as db:
        db.get(User, body["user"]["id"]).tier = "agency"
        db.commit()
    return {"Authorization": "Bearer " + body["access_token"]}, body["user"], email

def workspace(headers):
    result = client.post("/api/sessions", headers=headers, json={"business_goal": "Build a focused B2B service with evidence before spending."})
    assert result.status_code == 200, result.text
    return result.json()["id"]

def test_workspace_roles_and_optimistic_decision_versions():
    owner, _, _ = account(); stranger, _, email = account(); sid = workspace(owner)
    path = f"/api/studio/{sid}/records"
    payload = {"kind": "decision", "title": "Validate first", "body": "Interview buyers", "data": {"assumptions": "Demand unproven"}}
    assert client.get(path, headers=stranger).status_code == 404
    assert client.post(f"/api/studio/{sid}/members", headers=owner, json={"email": email, "role": "viewer"}).status_code == 200
    assert client.get(path, headers=stranger).status_code == 200
    assert client.post(path, headers=stranger, json=payload).status_code == 403
    r = client.post(path, headers=owner, json=payload).json()
    payload.update(version=r["version"], body="Interview ten buyers")
    assert client.put(path + "/" + r["id"], headers=owner, json=payload).status_code == 200
    assert client.put(path + "/" + r["id"], headers=owner, json=payload).status_code == 409
    versions = client.get(path, headers=owner).json()
    assert any(item["kind"] == "decision_revision" and item["body"] == "Interview buyers" for item in versions)

def test_financial_scenarios_calculate_real_inputs_and_reject_invalid_data():
    owner, _, _ = account(); sid = workspace(owner); path = f"/api/studio/{sid}/records"
    payload = {"kind": "scenario", "title": "Lean launch", "data": {"cash": 12000, "monthly_cost": 3000, "monthly_revenue": 1000, "ad_spend": 100, "customers": 10, "arpu": 50, "margin": .8, "churn": .1}}
    result = client.post(path, headers=owner, json=payload)
    assert result.status_code == 201, result.text
    assert result.json()["data"]["result"]["runway_months"] == 6
    assert result.json()["data"]["result"]["cac"] == 10
    payload["data"]["cash"] = -1
    assert client.post(path, headers=owner, json=payload).status_code == 422

def test_task_dependencies_and_assignment_enforce_workspace_boundaries():
    owner, user, _ = account(); sid = workspace(owner); path = f"/api/studio/{sid}/tasks"
    base = {"title": "Interview customers", "status": "Ready", "assignee_id": user["id"]}
    first = client.post(path, headers=owner, json=base).json()["id"]
    dependent = client.post(path, headers=owner, json={"title": "Ship experiment", "dependencies": [first]}).json()["id"]
    assert client.put(path + "/" + dependent, headers=owner, json={"title": "Ship experiment", "status": "Done", "dependencies": [first]}).status_code == 409
    assert client.put(path + "/" + first, headers=owner, json={"title": "Interview customers", "status": "Done"}).status_code == 200
    assert client.put(path + "/" + dependent, headers=owner, json={"title": "Ship experiment", "status": "Done", "dependencies": [first]}).status_code == 200
    assert client.post(path, headers=owner, json={"title": "Wrong dependency", "dependencies": ["missing"]}).status_code == 422

def test_knowledge_upload_search_and_ownership():
    owner, _, _ = account(); stranger, _, _ = account(); sid = workspace(owner)
    result = client.post(f"/api/studio/{sid}/knowledge", headers=owner, files={"file": ("notes.md", b"Our target buyer is a freelance accountant. They need invoice reconciliation.", "text/markdown")})
    assert result.status_code == 201, result.text
    search = client.get(f"/api/studio/{sid}/knowledge/search?q=invoice", headers=owner)
    assert search.json()["results"][0]["document_id"] == result.json()["id"]
    assert client.get(f"/api/studio/{sid}/knowledge/search?q=invoice", headers=stranger).status_code == 404
    assert client.post(f"/api/studio/{sid}/knowledge", headers=owner, files={"file": ("bad.exe", b"data", "application/octet-stream")}).status_code == 422

def test_canonical_run_streams_all_nine_and_persists_predictions():
    owner, _, _ = account(); sid = workspace(owner)
    result = client.post(f"/api/sessions/{sid}/messages", headers=owner, json={"content": "Assess a niche invoice automation service for accountants."})
    assert result.status_code == 200, result.text
    assert len(result.json()["reports"]) == 9
    predictions = client.get("/api/predictions", headers=owner).json()
    assert len(predictions["predictions"]) >= 9
    jobs = client.get("/api/jobs", headers=owner).json()
    assert jobs[0]["status"] == "done"
    assert len({r["agent"] for r in jobs[0]["reports"]}) == 9
    assert all(r["source"] == "local-template" for r in jobs[0]["reports"])

def test_revocation_and_password_change_invalidate_existing_tokens():
    owner, _, _ = account()
    active = client.get("/api/studio/security/sessions", headers=owner).json()
    assert client.delete("/api/studio/security/sessions/" + active[0]["id"], headers=owner).status_code == 200
    assert client.get("/api/auth/me", headers=owner).status_code == 401
    headers, _, _ = account()
    assert client.post("/api/account/password", headers=headers, json={"current_password": "correct-horse-battery", "new_password": "a-new-secure-password"}).status_code == 200
    assert client.get("/api/auth/me", headers=headers).status_code == 401

def test_board_review_is_completed_and_export_is_a_real_pdf():
    owner, _, _ = account(); sid = workspace(owner)
    assert client.post(f"/api/sessions/{sid}/board-meeting", headers=owner).status_code == 422
    client.post(f"/api/studio/{sid}/tasks", headers=owner, json={"title": "Interview buyers", "status": "Done"})
    review = client.post(f"/api/sessions/{sid}/board-meeting", headers=owner)
    assert review.status_code == 200, review.text
    assert "Progress" not in review.json()["title"]
    pdf = client.get(f"/api/studio/{sid}/reports/{review.json()['id']}/pdf", headers=owner)
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF-")
    from pypdf import PdfReader
    exported = PdfReader(io.BytesIO(pdf.content))
    assert review.json()["title"] in " ".join(page.extract_text() for page in exported.pages)

def test_account_deletion_cleans_new_studio_data():
    owner, _, _ = account(); sid = workspace(owner)
    assert client.post(f"/api/studio/{sid}/records", headers=owner, json={"kind": "decision", "title": "Deleteable", "data": {}}).status_code == 201
    result = client.request("DELETE", "/api/account", headers=owner, json={"password": "correct-horse-battery", "confirmation": "DELETE"})
    assert result.status_code == 200, result.text
    assert client.get("/api/auth/me", headers=owner).status_code == 401


def test_workspace_forecasts_and_analytics_are_scoped_and_respect_shared_roles():
    from datetime import datetime, timedelta
    from app.models import Prediction
    owner, _, _ = account(); editor, _, editor_email = account(); viewer, _, viewer_email = account(); stranger, _, _ = account()
    first = workspace(owner); second = workspace(owner)
    with SessionLocal() as db:
        alpha = Prediction(session_id=first, agent="CFO", statement="Alpha company only", confidence=65, due_at=datetime.utcnow() + timedelta(days=30))
        beta = Prediction(session_id=second, agent="CTO", statement="Beta company only", confidence=80, due_at=datetime.utcnow() + timedelta(days=30))
        db.add_all([alpha, beta]); db.commit(); alpha_id = alpha.id
    client.post(f"/api/studio/{first}/members", headers=owner, json={"email": editor_email, "role": "editor"})
    client.post(f"/api/studio/{first}/members", headers=owner, json={"email": viewer_email, "role": "viewer"})
    for actor in (owner, editor, viewer):
        forecasts = client.get(f"/api/predictions?session_id={first}", headers=actor)
        assert forecasts.status_code == 200 and [p["statement"] for p in forecasts.json()["predictions"]] == ["Alpha company only"]
        assert client.get(f"/api/analytics/overview?session_id={first}", headers=actor).json()["totals"]["predictions"] == 1
    for route in (f"/api/predictions?session_id={first}", f"/api/predictions/calibration?session_id={first}", f"/api/analytics/overview?session_id={first}"):
        assert client.get(route, headers=stranger).status_code == 404
    assert client.patch(f"/api/predictions/{alpha_id}", headers=viewer, json={"status": "hit"}).status_code == 403
    assert client.patch(f"/api/predictions/{alpha_id}", headers=editor, json={"status": "hit"}).status_code == 200
    assert client.get(f"/api/predictions/calibration?session_id={first}", headers=owner).json()["resolved_total"] == 1
    assert client.get(f"/api/predictions/calibration?session_id={second}", headers=owner).json()["resolved_total"] == 0


def test_cancelled_run_cannot_finalize_from_stale_cancellation_observation():
    import pytest
    from app.board_runner import execute_board_run
    from app.models import BusinessSession, Message, AgentReport, Prediction
    owner, user, _ = account(); sid = workspace(owner)
    with SessionLocal() as db:
        run = Job(user_id=user["id"], session_id=sid, prompt="Inspect launch evidence", status="running", progress_total=9)
        db.add(run); db.commit(); job_id = run.id
    checks = 0
    def stale_check():
        nonlocal checks
        checks += 1
        # Nine specialist chunks plus CEO synthesis precede the final check.
        if checks == 11:
            with SessionLocal() as cancel:
                cancel.query(Job).filter_by(id=job_id).update({"status": "cancelled"})
                cancel.commit()
        return False
    with SessionLocal() as db:
        with pytest.raises(InterruptedError):
            execute_board_run(db, db.get(BusinessSession, sid), "Inspect launch evidence", on_event=lambda event: None, cancelled=stale_check, job_id=job_id)
    with SessionLocal() as db:
        assert db.get(Job, job_id).status == "cancelled"
        assert db.query(Message).filter_by(session_id=sid, role="assistant").count() == 0
        assert db.query(AgentReport).filter_by(session_id=sid).count() == 0
        assert db.query(Prediction).filter_by(session_id=sid).count() == 0


def test_durable_cancellation_retry_and_recovery_preserve_one_canonical_result(monkeypatch):
    from threading import Event
    from app import jobs
    from app.models import Message
    owner, user, _ = account(); stranger, _, _ = account(); sid = workspace(owner)
    entered, release, left = Event(), Event(), Event()
    original = jobs.execute_board_run
    def delayed(*args, **kwargs):
        entered.set()
        try:
            if not release.wait(10):
                raise RuntimeError("Isolated worker did not receive its test release.")
            return original(*args, **kwargs)
        finally:
            left.set()
    monkeypatch.setattr(jobs, "execute_board_run", delayed)
    initial = client.post(f"/api/jobs/board-run/{sid}", headers=owner, json={"content": "Test durable cancellation."}).json()["job_id"]
    try:
        assert entered.wait(5)
        assert client.post(f"/api/jobs/{initial}/cancel", headers=stranger).status_code == 404
        assert client.post(f"/api/jobs/{initial}/cancel", headers=owner).json()["status"] == "cancelled"
    finally:
        release.set()
    assert left.wait(10)
    monkeypatch.setattr(jobs, "execute_board_run", original)
    retry = client.post(f"/api/jobs/{initial}/retry", headers=owner)
    assert retry.status_code == 200 and retry.json()["job_id"] != initial
    retried_id = retry.json()["job_id"]
    deadline = time.monotonic() + 15
    state = {}
    while time.monotonic() < deadline:
        state = client.get(f"/api/jobs/{retried_id}", headers=owner).json()
        if state.get("status") in {"done", "failed"}:
            break
        time.sleep(.05)
    assert state["status"] == "done" and len(state["reports"]) == 9
    assert client.post(f"/api/jobs/{retried_id}/retry", headers=owner).status_code == 409
    assert client.post(f"/api/jobs/{retried_id}/cancel", headers=owner).json()["status"] == "done"
    with SessionLocal() as db:
        assert db.query(Message).filter_by(session_id=sid, role="assistant").count() == 1
        assert db.query(RunEvent).filter_by(job_id=retried_id).count() == 10
