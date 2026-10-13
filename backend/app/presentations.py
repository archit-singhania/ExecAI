"""Source-snapshotted board decks, real exports and existing revocable sharing."""
import io
import json
import math
import textwrap
from typing import Literal
from xml.sax.saxutils import escape
from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel, Field
from sqlalchemy import update
from sqlalchemy.orm import Session
from app.access import workspace_access
from app.auth import get_current_user
from app.database import get_db
from app.models import AgentReport, BusinessSession, Task, User
from app.studio_models import StudioRecord
from app.studio import record_out, record_activity
from app.share import create_limit

router = APIRouter(prefix="/api/studio", tags=["board-presentations"])

def export_slides(slides):
    """Bound text per page rather than clipping long editable slide bodies."""
    for source in slides:
        lines = [part for line in source["body"].split("\n") for part in (textwrap.wrap(line, width=85) or [""])]
        chunks = [lines[i:i + 18] for i in range(0, len(lines), 18)] or [[]]
        for index, chunk in enumerate(chunks):
            yield {**source, "title": source["title"] + (f" · continued {index + 1}" if index else ""), "body": "\n".join(chunk), "notes": source["notes"] if index == 0 else ""}

class GenerateIn(BaseModel):
    decision_id: str = Field(min_length=1, max_length=36)
    scenario_ids: list[str] = Field(default_factory=list, max_length=6)

class SlideIn(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    body: str = Field(default="", max_length=2000)
    notes: str = Field(default="", max_length=5000)

class DeckIn(BaseModel):
    title: str = Field(min_length=1, max_length=240)
    slides: list[SlideIn] = Field(min_length=1, max_length=20)
    version: int

def owned(db, sid, rid):
    row = db.get(StudioRecord, rid)
    if not row or row.session_id != sid or row.kind != "presentation":
        raise HTTPException(404, "Presentation not found.")
    return row

@router.get("/{session_id}/presentations")
def list_decks(session_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    workspace_access(db, session_id, user)
    return [record_out(r) for r in db.query(StudioRecord).filter_by(session_id=session_id, kind="presentation").order_by(StudioRecord.updated_at.desc()).all()]

@router.post("/{session_id}/presentations", status_code=201)
def generate(session_id: str, payload: GenerateIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    workspace_access(db, session_id, user, write=True)
    decision = db.get(StudioRecord, payload.decision_id)
    if not decision or decision.session_id != session_id or decision.kind != "decision":
        raise HTTPException(404, "Decision not found.")
    from app.freshness import latest_baselines
    baseline = next((r for r in latest_baselines(db, session_id) if r.title == decision.id), None)
    evidence = json.loads(baseline.data)["evidence"] if baseline else []
    context = json.loads(decision.data)
    slides = [{"title": decision.title, "body": decision.body[:2000], "notes": "Review the recommendation before presenting."},
              {"title": "Alternatives and assumptions", "body": f"Alternatives: {context.get('alternatives', 'Not recorded')}\nAssumptions: {context.get('assumptions', 'Not recorded')}"[:2000], "notes": "Separate documented evidence from assumptions."},
              {"title": "Supporting evidence", "body": "\n".join(f"[{i + 1}] {r['title']} · revision {r['version']}" for i, r in enumerate(evidence))[:2000] or "No approved evidence baseline. Add and review supporting evidence before distribution.", "notes": "Source details are available in the saved approval baseline."}]
    scenarios = []
    for rid in dict.fromkeys(payload.scenario_ids):
        scenario = db.get(StudioRecord, rid)
        if not scenario or scenario.session_id != session_id or scenario.kind != "scenario":
            raise HTTPException(404, "Scenario not found in this workspace.")
        value = json.loads(scenario.data); scenarios.append({"title": scenario.title, "version": scenario.version, "data": value})
    chart = None
    if scenarios:
        amounts = [s["data"].get("monthly_revenue") for s in scenarios]
        if all(isinstance(v, (float, int)) and not isinstance(v, bool) and math.isfinite(v) for v in amounts):
            chart = {"labels": [s["title"] for s in scenarios], "values": amounts, "label": "Entered monthly revenue (project currency)"}
        slides.append({"title": "Scenario comparison", "body": "\n".join(s["title"] + ": " + json.dumps(s["data"], ensure_ascii=False) for s in scenarios)[:2000], "notes": "Entered assumptions; not forecast guarantees."})
    tasks = db.query(Task).filter_by(session_id=session_id).order_by(Task.created_at.desc()).limit(10).all()
    slides.append({"title": "Actions and review", "body": "\n".join(f"{t.title} · {t.status}" for t in tasks)[:2000] or "No workspace tasks recorded. Agree owners and review dates.", "notes": "These are workspace tasks, not inferred commitments."})
    row = StudioRecord(session_id=session_id, author_id=user.id, kind="presentation", title=decision.title,
        body="Editable board presentation", data=json.dumps({"decision_id": decision.id, "decision_version": decision.version, "baseline_id": baseline.id if baseline else None, "sources": evidence, "scenarios": scenarios, "chart": chart, "slides": slides}))
    db.add(row); record_activity(db, session_id, user, "Board presentation generated", row.title)
    db.commit(); db.refresh(row)
    return record_out(row)

@router.post("/{session_id}/presentations/{record_id}/share")
def publish(session_id: str, record_id: str, db: Session = Depends(get_db), user: User = Depends(create_limit)):
    workspace = workspace_access(db, session_id, user, write=True)
    if workspace.user_id != user.id:
        raise HTTPException(403, "Only the workspace owner can publish a presentation.")
    row = owned(db, session_id, record_id); value = json.loads(row.data)
    source = f"presentation:{row.id}:v{row.version}"
    report = db.query(AgentReport).filter_by(session_id=session_id, report_type="presentation", source=source).first()
    if not report:
        public = {"slides": [{"title": s["title"], "body": s["body"]} for s in value["slides"]], "chart": value.get("chart"), "version": row.version}
        report = AgentReport(session_id=session_id, agent="CEO.ai presentation", report_type="presentation", title=row.title,
                             summary="User-reviewed board presentation snapshot", bullets=json.dumps(public), source=source)
        db.add(report); db.commit(); db.refresh(report)
    from app.share import create_share_link
    result = create_share_link(report.id, db, user)
    from app.config import get_settings
    return {"slug": result["slug"], "url": get_settings().app_base_url + "/p/" + result["slug"], "version": row.version}

@router.delete("/{session_id}/presentations/{record_id}/share")
def revoke(session_id: str, record_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    workspace = workspace_access(db, session_id, user, write=True)
    if workspace.user_id != user.id:
        raise HTTPException(403, "Only the workspace owner can revoke presentation links.")
    row = owned(db, session_id, record_id)
    reports = db.query(AgentReport).filter(AgentReport.session_id == session_id, AgentReport.source.like(f"presentation:{row.id}:%")).all()
    from app.share import revoke_share_link
    for report in reports:
        revoke_share_link(report.id, db, user)
    return {"revoked_snapshots": len(reports)}

@router.put("/{session_id}/presentations/{record_id}")
def save(session_id: str, record_id: str, payload: DeckIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    workspace_access(db, session_id, user, write=True)
    db.execute(update(BusinessSession).where(BusinessSession.id == session_id).values(title=BusinessSession.title))
    row = owned(db, session_id, record_id)
    if row.version != payload.version:
        raise HTTPException(409, "Presentation changed. Reload before saving.")
    if not payload.title.strip() or any(not s.title.strip() for s in payload.slides):
        raise HTTPException(422, "Presentation and slide titles cannot be blank.")
    value = json.loads(row.data); value["slides"] = [s.model_dump() for s in payload.slides]
    row.title, row.data, row.version = payload.title.strip(), json.dumps(value), row.version + 1
    db.commit(); db.refresh(row)
    return record_out(row)

@router.get("/{session_id}/presentations/{record_id}/export/{format}")
def export(session_id: str, record_id: str, format: Literal["pdf", "pptx"], include_notes: bool = False, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    workspace_access(db, session_id, user)
    from app.plans import get_plan
    if not get_plan(user.tier).exports:
        raise HTTPException(402, "Presentation exports are available on Pro and above.")
    row = owned(db, session_id, record_id); value = json.loads(row.data); output = io.BytesIO()
    if format == "pdf":
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.colors import HexColor
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        styles = getSampleStyleSheet(); flow = []
        styles["BodyText"].fontSize = 13; styles["BodyText"].leading = 18
        styles["Title"].textColor = HexColor("#19273d")
        styles["BodyText"].textColor = HexColor("#39465d")
        for slide in export_slides(value["slides"]):
            if flow: flow.append(PageBreak())
            flow.extend([Paragraph(escape(slide["title"]), styles["Title"]), Spacer(1, 20), Paragraph(escape(slide["body"]).replace("\n", "<br/>"), styles["BodyText"])])
            if include_notes: flow.extend([Spacer(1, 20), Paragraph("Speaker notes: " + escape(slide["notes"]), styles["Normal"])])
        chart = value.get("chart")
        if chart:
            from reportlab.graphics.shapes import Drawing
            from reportlab.graphics.charts.barcharts import VerticalBarChart
            drawing = Drawing(650, 250); bars = VerticalBarChart()
            bars.x, bars.y, bars.width, bars.height = 60, 65, 550, 160
            bars.data = [chart["values"]]; bars.categoryAxis.categoryNames = [s[:25] for s in chart["labels"]]
            bars.bars[0].fillColor = HexColor("#4968b5")
            bars.valueAxis.valueMin = min(0, min(chart["values"])); bars.valueAxis.valueMax = max(1, max(chart["values"]))
            drawing.add(bars)
            flow.extend([PageBreak(), Paragraph(escape(chart["label"]), styles["Title"]), Spacer(1, 20), drawing])
        SimpleDocTemplate(output, pagesize=landscape(A4), title=row.title, author="CEO.ai").build(flow)
        media = "application/pdf"
    else:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5)
        for source in export_slides(value["slides"]):
            slide = deck.slides.add_slide(deck.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(23, 31, 46)
            for content, y, size in [(source["title"], .55, 24), (source["body"], 2.05, 16)]:
                box = slide.shapes.add_textbox(Inches(.8), Inches(y), Inches(11.7), Inches(4.9 if y > 1 else 1.3))
                box.text_frame.word_wrap = True; box.text_frame.text = content
                for paragraph in box.text_frame.paragraphs:
                    paragraph.font.size = Pt(size); paragraph.font.color.rgb = RGBColor(245, 246, 250)
            if include_notes: slide.notes_slide.notes_text_frame.text = source["notes"]
        chart = value.get("chart")
        if chart:
            from pptx.chart.data import CategoryChartData
            from pptx.enum.chart import XL_CHART_TYPE
            slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = chart["label"]
            data = CategoryChartData(); data.categories = chart["labels"]; data.add_series("Entered revenue", chart["values"])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(11), Inches(5), data)
        deck.save(output); media = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    return Response(output.getvalue(), media_type=media, headers={"Content-Disposition": f'attachment; filename="ceoai-board-{row.id}.{format}"'})
